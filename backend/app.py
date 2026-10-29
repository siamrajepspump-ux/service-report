"""

Siamraj Service Report - Unified Backend API Server

"""



import os

import sys

import csv

import json

import uuid

import shutil

import time

from typing import Dict, Any, List, Optional, Tuple

from datetime import datetime

import openpyxl

from fastapi import FastAPI, UploadFile, File, Form, Query, HTTPException, BackgroundTasks, Body

from fastapi.middleware.cors import CORSMiddleware

from fastapi.staticfiles import StaticFiles

from fastapi.responses import FileResponse, JSONResponse

from pydantic import BaseModel



# Add current directory and pdf_engine to sys.path

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))

BASE_DIR = os.path.dirname(CURRENT_DIR)

sys.path.append(CURRENT_DIR)

sys.path.append(os.path.join(CURRENT_DIR, "pdf_engine"))



from iso_engine import IsoPdfEngine
import gas_sync



app = FastAPI(

    title="Siamraj Service Report API",

    version="2.0.0",

    description="Unified API for Service Report, Inspection, Photo Report & ISO PDF Automation"

)



# Enable CORS for local development and mobile network access

app.add_middleware(

    CORSMiddleware,

    allow_origins=["*"],

    allow_credentials=True,

    allow_methods=["*"],

    allow_headers=["*"],

)



@app.middleware("http")

async def add_no_cache_header(request, call_next):

    response = await call_next(request)

    if any(request.url.path.endswith(ext) for ext in [".js", ".html", ".css", ".json"]):

        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"

        response.headers["Pragma"] = "no-cache"

        response.headers["Expires"] = "0"

    return response



DATA_DIR = os.path.join(BASE_DIR, "data")

STORAGE_DIR = os.path.join(BASE_DIR, "storage")

REPORTS_DIR = os.path.join(STORAGE_DIR, "reports")

PHOTOS_DIR = os.path.join(STORAGE_DIR, "photos")

OUTPUT_PDF_DIR = os.path.join(STORAGE_DIR, "output_pdf")

FRONTEND_DIR = os.path.join(BASE_DIR, "frontend", "public")



for d in [DATA_DIR, STORAGE_DIR, REPORTS_DIR, PHOTOS_DIR, OUTPUT_PDF_DIR]:

    os.makedirs(d, exist_ok=True)



pdf_engine = IsoPdfEngine()



# Primary source on Google Drive and local fallback

PRIMARY_TICKETS_CSV = r"G:\My Drive\Service_App_Data\Tickets_Cleaned.csv"

FALLBACK_TICKETS_CSV = os.path.join(DATA_DIR, "crm_tickets_cache.csv")



# In-memory CRM cache

crm_cache: List[Dict[str, Any]] = []

crm_cache_mtime: float = 0.0



def load_crm_cache(force: bool = False) -> List[Dict[str, Any]]:

    global crm_cache, crm_cache_mtime



    # Auto clean if raw Tickets.csv is newer than Tickets_Cleaned.csv

    raw_csv = r"G:\My Drive\Service_App_Data\Tickets.csv"

    clean_csv = r"G:\My Drive\Service_App_Data\Tickets_Cleaned.csv"

    if os.path.exists(raw_csv) and os.path.exists(clean_csv):

        try:

            if os.path.getmtime(raw_csv) > os.path.getmtime(clean_csv):

                import sys

                app_data_dir = r"G:\My Drive\Service_App_Data"

                if app_data_dir not in sys.path:

                    sys.path.insert(0, app_data_dir)

                from clean_data import run_data_cleaning

                run_data_cleaning()

                print("[CRM Cache] Auto-cleaned newer Tickets.csv on the fly!")

        except Exception as ex:

            print(f"[CRM Cache] Auto-clean notice: {ex}")



    target_path = PRIMARY_TICKETS_CSV if os.path.exists(PRIMARY_TICKETS_CSV) else FALLBACK_TICKETS_CSV

    if not os.path.exists(target_path):

        print(f"[CRM Cache] Neither primary nor fallback CSV exists.")

        return crm_cache



    try:

        current_mtime = os.path.getmtime(target_path)

        if not force and current_mtime == crm_cache_mtime and crm_cache:

            return crm_cache



        is_service_app_data = (target_path == PRIMARY_TICKETS_CSV)

        with open(target_path, mode="r", encoding="utf-8-sig") as f:

            reader = csv.DictReader(f)

            loaded_items = []

            for r in reader:

                if is_service_app_data:

                    t_no = str(r.get("Title") or r.get("Ticket Number") or "").strip()

                    if not t_no or t_no.lower() == "nan":

                        continue

                    clean_str = lambda val: "" if (val is None or str(val).strip().lower() in ["nan", "none"]) else str(val).strip()

                    item = {k: clean_str(v) for k, v in r.items()}

                    item["TicketNo"] = t_no

                    item["TicketKey"] = clean_str(r.get("Ticket Number"))

                    item["Customer"] = clean_str(r.get("Organization Name"))

                    item["Site"] = clean_str(r.get("Address or Site"))

                    item["ContactPerson"] = clean_str(r.get("Site Contact Person"))

                    item["ContactPhone"] = clean_str(r.get("Site Contact Phone"))

                    item["Model"] = clean_str(r.get("Model"))

                    item["SerialNo"] = clean_str(r.get("Serial Number"))

                    item["Product"] = clean_str(r.get("Product"))

                    item["ServiceBy"] = clean_str(r.get("Service By"))

                    item["SalesOrder"] = clean_str(r.get("Sales Orders"))

                    item["Description"] = clean_str(r.get("Description"))

                    item["ProblemDetail"] = clean_str(r.get("Description"))

                    item["Cause"] = clean_str(r.get("Cause"))

                    item["Prevention"] = clean_str(r.get("Prevention"))

                    item["Solution"] = clean_str(r.get("Solution"))

                    item["AppointmentDate"] = clean_str(r.get("Appointment Date"))

                    item["CompletedDate"] = clean_str(r.get("Completed Date"))

                    item["CreatedDate"] = clean_str(r.get("Created Time"))

                    item["Status"] = clean_str(r.get("Status"))

                    loaded_items.append(item)

                else:

                    loaded_items.append(r)



            crm_cache = loaded_items

            crm_cache_mtime = current_mtime

            print(f"[CRM Cache] Loaded {len(crm_cache)} tickets from {target_path} (updated: {time.ctime(current_mtime)})")

    except Exception as e:

        print(f"[CRM Cache] Error loading {target_path}: {e}")

    return crm_cache



load_crm_cache(force=True)



# =============================================================================

# API ROUTES: TICKETS & CRM

# =============================================================================

@app.get("/api/health")

def health_check():

    load_crm_cache()

    return {

        "status": "online",

        "timestamp": datetime.now().isoformat(),

        "tickets_cached": len(crm_cache),

        "source": PRIMARY_TICKETS_CSV if os.path.exists(PRIMARY_TICKETS_CSV) else FALLBACK_TICKETS_CSV,

        "version": "2.0.0"

    }



@app.get("/api/tickets")

def list_tickets(q: Optional[str] = Query(None), page: int = 1, limit: int = 25, sort: str = "desc"):

    """Search and paginate cached CRM tickets directly from Google Drive"""

    load_crm_cache()

    results = crm_cache

    if q:

        query = q.lower().strip()

        results = [

            t for t in crm_cache

            if query in str(t.get("TicketNo", "")).lower()

            or query in str(t.get("TicketKey", "")).lower()

            or query in str(t.get("Customer", "")).lower()

            or query in str(t.get("Site", "")).lower()

            or query in str(t.get("Model", "")).lower()

            or query in str(t.get("SerialNo", "")).lower()

            or query in str(t.get("ServiceBy", "")).lower()

            or query in str(t.get("SalesOrder", "")).lower()

        ]

        results.sort(key=lambda t: 0 if str(t.get("TicketNo", "")).lower() == query else (1 if str(t.get("TicketKey", "")).lower() == query else 2))

    elif sort == "desc":

        results = list(reversed(crm_cache))



    total = len(results)

    start = (page - 1) * limit

    end = start + limit

    items = results[start:end]

    return {

        "items": items,

        "total": total,

        "page": page,

        "limit": limit

    }



@app.get("/api/tickets/{ticket_no}")

def get_ticket(ticket_no: str):

    """Retrieve full details of a specific ticket by Request No. or Ticket Key"""

    load_crm_cache()

    t_clean = ticket_no.strip().lower()

    for t in crm_cache:

        t_no = str(t.get("TicketNo", "")).strip().lower()

        t_key = str(t.get("TicketKey", "")).strip().lower()

        if t_no == t_clean or t_key == t_clean:

            return t

    raise HTTPException(status_code=404, detail="Ticket not found")



@app.post("/api/tickets/reload")

def reload_tickets():

    """Force reload tickets from Google Drive"""

    tickets = load_crm_cache(force=True)

    target_path = PRIMARY_TICKETS_CSV if os.path.exists(PRIMARY_TICKETS_CSV) else FALLBACK_TICKETS_CSV

    return {

        "status": "success",

        "tickets_cached": len(tickets),

        "source": target_path,

        "timestamp": datetime.now().isoformat()

    }



# =============================================================================

# API ROUTES: BOM (VIKING & GOULDS) & INSPECTION SPECS

# =============================================================================

def _load_insp_csv_data():

    insp_csv_path = os.path.join(DATA_DIR, "bom", "viking", "insp.csv")

    if not os.path.exists(insp_csv_path):

        return []

    try:

        with open(insp_csv_path, "r", encoding="utf-8-sig") as f:

            return list(csv.DictReader(f))

    except:

        return []



def _resolve_pir_file(model_str: str, insp_files: set, models_data: dict) -> Optional[str]:

    m = model_str.strip().upper()

    if not m:

        return None

    if m in insp_files:

        return m

    if f"PIR-{m}" in insp_files:

        return f"PIR-{m}"



    m_info = models_data.get(m, {})

    if m_info.get("insp") and m_info["insp"] in insp_files:

        return m_info["insp"]

    for opt in m_info.get("inspOptions", []):

        if opt in insp_files:

            return opt



    sizes = ["ROTOR", "CASING", "BUSHING", "SHAFT", "IDLER", "AL", "AK", "AS", "GG", "HJ", "HL", "FH", "KK", "LQ", "LL", "LS", "QS", "RS", "KE", "KKE", "ALE", "ATE", "HLE", "LSE", "C", "F", "G", "H", "J", "K", "L", "M", "N", "P", "Q", "R", "W"]

    sizes.sort(key=len, reverse=True)



    matched_size = None

    series_part = ""

    for s in sizes:

        if m.startswith(s):

            matched_size = s

            series_part = m[len(s):]

            break



    clean_series = series_part

    alt_series = series_part

    if series_part.startswith("4") and len(series_part) > 2:

        if series_part.startswith("41") and len(series_part) >= 4 and series_part[2] in "1234567890":

            alt_series = series_part[1:]

        elif series_part.startswith("4") and len(series_part) == 3 and series_part[1:] == "32":

            alt_series = "32"



    candidates = [

        f"PIR-{series_part}-{matched_size}",

        f"PIR-{alt_series}-{matched_size}",

        f"PIR-{matched_size}-{series_part}",

        f"PIR-{matched_size}{series_part}",

        f"{matched_size}{series_part}",

        f"{m}"

    ]

    if matched_size == "LQ":

        candidates.extend([f"PIR-{series_part}-LLQ", f"PIR-{alt_series}-LLQ"])



    for c in candidates:

        if c in insp_files:

            return c



    for f in insp_files:

        if matched_size and f"-{matched_size}" in f:

            if series_part in f or alt_series in f:

                return f



    return None



def _parse_insp_num(s: str) -> Optional[float]:

    if not s:

        return None

    s = s.strip()

    if s.startswith("."):

        s = "0" + s

    if s.count(".") > 1:

        p = s.split(".")

        s = p[0] + "." + "".join(p[1:])

    try:

        return float(s)

    except:

        return None



def _build_inspection_19_rows(model_str: str) -> Tuple[Optional[str], List[Dict[str, Any]]]:

    insp_rows = _load_insp_csv_data()

    if not insp_rows:

        return None, []

    insp_files = set(r["File"] for r in insp_rows)



    models_file = os.path.join(DATA_DIR, "bom", "viking", "models.json")

    models_data = {}

    if os.path.exists(models_file):

        try:

            with open(models_file, "r", encoding="utf-8") as f:

                models_data = json.load(f).get("models", {})

        except:

            pass



    pir_file = _resolve_pir_file(model_str, insp_files, models_data)

    if not pir_file:

        return None, []



    items = [r for r in insp_rows if r["File"] == pir_file]

    sections = [

        ("ROTOR OD", ["ROTOR O.D.", "ROTOR OD"], "U/L"),

        ("ROTOR ID", ["ROTOR I.D.", "ROTOR ID"], "U/L"),

        ("ROTOR TOOTH LENGTH", ["ROTOR TOOTH LENGTH"], "-"),

        ("IDLER OD", ["IDLER O.D.", "IDLER OD"], "U/L"),

        ("IDLER BUSHING ID", ["IDLER (BUSHING) I.D.", "IDLER BUSHING I.D.", "BUSHING I.D."], "U/L"),

        ("IDLER TOOTH LENGTH", ["IDLER TOOTH LENGTH"], "-"),

        ("IDLER PIN OD", ["IDLER PIN O.D.", "IDLER PIN O.D. (1)"], "U/L"),

        ("SHAFT OD", ["SHAFT O.D.", "SHAFT O.D. (CASING)", "SHAFT O.D. (HEAD)"], "U/L"),

        ("SHAFT BUSHING ID", ["SHAFT BUSHING I.D.", "ADAPTOR BUSHING I.D.", "CANISTER BUSHING I.D."], "U/L"),

        ("CRESCENT LENGTH", ["CRESCENT LENGTH"], "-"),

        ("CASING ID", ["CASING I.D.", "CASING ROTOR CHAMBER"], "U/L")

    ]



    rows = []

    for title, match_keys, pos_type in sections:

        found = None

        for it in items:

            if it.get("Description") in match_keys:

                found = it

                break



        cmp_op = (found.get("Cmp") or "").strip() if found else ""

        std_raw = (found.get("Standard") or "").strip() if found else ""

        val_in = _parse_insp_num(std_raw)

        val_mm = round(val_in * 25.4, 2) if val_in is not None else None



        std_in_str = f"{cmp_op} {val_in:.4f}\"" if val_in is not None else (f"{cmp_op} {std_raw}" if std_raw else "-")

        std_mm_str = f"{cmp_op} {val_mm:.2f} mm" if val_mm is not None else "-"

        std_min = f"{val_mm:.2f}" if (cmp_op == "≥" and val_mm is not None) else ""

        std_max = f"{val_mm:.2f}" if (cmp_op == "≤" and val_mm is not None) else ""



        # Initial value for inputs: default to '-' as requested by technician

        if pos_type == "U/L":

            rows.append({

                "part": title, "pos": "UPPER", "cmp": cmp_op,

                "std_in": std_in_str, "std_mm": std_mm_str,

                "min": std_min, "max": std_max, "val_mm": val_mm,

                "in_a": "-", "in_b": "-",

                "final_a": "-", "final_b": "-"

            })

            rows.append({

                "part": title, "pos": "LOWER", "cmp": cmp_op,

                "std_in": std_in_str, "std_mm": std_mm_str,

                "min": std_min, "max": std_max, "val_mm": val_mm,

                "in_a": "-", "in_b": "-",

                "final_a": "-", "final_b": "-"

            })

        else:

            rows.append({

                "part": title, "pos": "-", "cmp": cmp_op,

                "std_in": std_in_str, "std_mm": std_mm_str,

                "min": std_min, "max": std_max, "val_mm": val_mm,

                "in_a": "-", "in_b": "-",

                "final_a": "-", "final_b": "-"

            })



    return pir_file, rows



@app.get("/api/bom/inspection/models")

def get_inspection_models():

    models_file = os.path.join(DATA_DIR, "bom", "viking", "models.json")

    insp_csv_path = os.path.join(DATA_DIR, "bom", "viking", "insp.csv")

    

    unique_models = set()

    if os.path.exists(models_file):

        try:

            with open(models_file, "r", encoding="utf-8") as f:

                data = json.load(f).get("models", {})

                for k in data.keys():

                    unique_models.add(k)

        except:

            pass



    insp_files = []

    if os.path.exists(insp_csv_path):

        try:

            with open(insp_csv_path, "r", encoding="utf-8-sig") as f:

                r = csv.DictReader(f)

                for row in r:

                    unique_models.add(row["File"])

                    insp_files.append(row["File"])

        except:

            pass



    # Curate popular models list

    popular = [

        "H4124A", "HL4124A", "K4124A", "KK4124A", "L4124A", "LQ4124A", "LL4124A", "LS4124A", "Q4124A", "QS4124A", "M4124A",

        "H124A", "HL124A", "K124A", "KK124A", "L124A", "LQ124A", "LL124A", "LS124A", "Q124A", "QS124A", "M124A",

        "H4123A", "HL4123A", "K4123A", "KK4123A", "L4123A", "LQ4123A", "LL4123A", "LS4123A", "Q4123A", "QS4123A",

        "H123A", "HL123A", "K123A", "KK123A", "L123A", "LQ123A", "LL123A", "LS123A", "Q123A", "QS123A",

        "H4127A", "HL4127A", "K4127A", "KK4127A", "L4127A", "LQ4127A", "LL4127A", "LS4127A", "Q4127A", "QS4127A",

        "H127A", "HL127A", "K127A", "KK127A", "L127A", "LQ127A", "LL127A", "LS127A", "Q127A", "QS127A",

        "H32", "HL32", "K32", "KK32", "L32", "LQ32", "LL32", "Q32", "M32", "N32",

        "H432", "HL432", "AK4195", "AL4195", "GG4195", "HJ4195", "HL4195", "AS4195",

        "AK4197", "AL4197", "GG4197", "HJ4197", "HL4197", "AS4197",

        "KS157B", "157B"

    ]

    all_sorted = sorted(list(unique_models))

    return {

        "popular": popular,

        "all_models": all_sorted,

        "pir_files": sorted(list(set(insp_files)))

    }



@app.get("/api/bom/inspection/specs")

def get_inspection_specs(model: str = Query(...)):

    pir_file, rows = _build_inspection_19_rows(model)

    if not pir_file or not rows:

        return {

            "ok": False,

            "model": model,

            "detail": f"ไม่พบข้อมูลเกณฑ์มาตรฐานสำหรับรุ่น '{model}' ในฐานข้อมูล Inspection",

            "rows": []

        }

    return {

        "ok": True,

        "model": model,

        "pir_file": pir_file,

        "rows": rows

    }



@app.get("/api/bom/viking/models")

def get_viking_models():

    models_file = os.path.join(DATA_DIR, "bom", "viking", "models.json")

    if os.path.exists(models_file):

        try:

            with open(models_file, "r", encoding="utf-8") as f:

                return json.load(f)

        except Exception as e:

            raise HTTPException(status_code=500, detail=str(e))

    return []



@app.get("/api/bom/viking/specs")

def get_viking_specs():

    stock_file = os.path.join(DATA_DIR, "bom", "viking", "stock.json")

    stock_data = {}

    if os.path.exists(stock_file):

        try:

            with open(stock_file, "r", encoding="utf-8") as f:

                stock_data = json.load(f)

        except:

            pass

    return {

        "stock": stock_data

    }



@app.get("/api/bom/goulds/groups")

def get_goulds_groups():

    return [

        {"id": "STi", "name": "Goulds 3196 STi Frame", "sizes": ["1.5x3-6", "1.5x3-8", "2x3-8", "3x4-8", "3x4-8G"]},

        {"id": "MTi", "name": "Goulds 3196 MTi Frame", "sizes": ["1x2-10", "1.5x3-10", "2x3-10", "3x4-10", "4x6-10", "2x3-13", "3x4-13"]},

        {"id": "LTi", "name": "Goulds 3196 LTi Frame", "sizes": ["4x6-13", "6x8-13", "8x10-13", "4x6-15", "6x8-15", "8x10-15"]},

        {"id": "XLT-i", "name": "Goulds 3196 XLT-i Frame", "sizes": ["6x8-15G", "8x10-15G", "8x10-16G", "10x12-16"]},

        {"id": "i-17", "name": "Goulds 3196 i-17 Frame", "sizes": ["10x12-17", "12x14-17"]}

    ]



SHARED_VIKING_APP = r"g:\Shared drives\Service Report\BOM_Viking Part List\_app"

SHARED_GOULDS_DIR = r"g:\Shared drives\Service Report\BOM_Goulds Part List"



def _load_viking_bom_parts(model_query: str) -> Tuple[Optional[Dict[str, Any]], Optional[str]]:

    models_file = os.path.join(DATA_DIR, "bom", "viking", "models.json")

    if not os.path.exists(models_file):

        models_file = os.path.join(SHARED_VIKING_APP, "models.json")

    if not os.path.exists(models_file):

        return None, "models.json not found"

        

    with open(models_file, "r", encoding="utf-8") as f:

        models_dict = json.load(f).get("models", {})

        

    mq_upper = model_query.strip().upper()

    matched_key = None

    if mq_upper in models_dict:

        matched_key = mq_upper

    else:

        for k in models_dict:

            if k.upper() == mq_upper:

                matched_key = k

                break

        if not matched_key:

            for k in models_dict:

                if k.upper().startswith(mq_upper) or mq_upper.startswith(k.upper()):

                    matched_key = k

                    break

        if not matched_key:

            for k in models_dict:

                if mq_upper in k.upper():

                    matched_key = k

                    break

                    

    if not matched_key:

        return None, f"ไม่พบข้อมูล BOM ของรุ่น '{model_query}' ในฐานข้อมูล Viking"

        

    info = models_dict[matched_key]

    parts_key = info.get("parts")

    diagram = info.get("diagram")

    

    bom_dirs = [

        os.path.join(DATA_DIR, "bom", "viking", "bom"),

        os.path.join(SHARED_VIKING_APP, "bom")

    ]

    

    parts_file = None

    for bd in bom_dirs:

        cand = os.path.join(bd, f"{parts_key}.json")

        if os.path.exists(cand):

            parts_file = cand

            break

            

    if not parts_file:

        return None, f"ไม่พบไฟล์ BOM ({parts_key}.json) ของรุ่น {matched_key}"

        

    with open(parts_file, "r", encoding="utf-8") as f:

        bom_data = json.load(f)

        

    stock_file = os.path.join(DATA_DIR, "bom", "viking", "stock.json")

    stock_dict = {}

    if os.path.exists(stock_file):

        try:

            with open(stock_file, "r", encoding="utf-8") as f:

                stock_dict = json.load(f).get("parts", {})

        except:

            pass

            

    res_parts = []

    for p in bom_data.get("parts", []):

        pno = str(p.get("partNo") or "").strip()

        stock_info = stock_dict.get(pno, {})

        res_parts.append({

            "item": str(p.get("item") or "").strip(),

            "part_no": pno,

            "desc": p.get("desc") or "",

            "material": p.get("material") or "",

            "qty": p.get("qty") or 1,

            "price": stock_info.get("unit"),

            "stock_qty": stock_info.get("qty"),

            "stock_status": stock_info.get("status")

        })

        

    return {

        "ok": True,

        "brand": "VIKING",

        "model": matched_key,

        "title": bom_data.get("title") or matched_key,

        "diagram": f"/api/bom/diagram/{diagram}" if diagram else None,

        "count": len(res_parts),

        "parts": res_parts

    }, None



def _load_goulds_bom_parts(model_query: str) -> Tuple[Optional[Dict[str, Any]], Optional[str]]:

    goulds_dirs = [

        os.path.join(DATA_DIR, "bom", "goulds"),

        SHARED_GOULDS_DIR

    ]

    mq_clean = model_query.upper().replace("GOULDS", "").replace("3196", "").replace(" ", "").strip()

    

    file_map = {

        "STI": ("Goulds_3196_PartList_STi.xlsx", "STi", "ผังรวม_3196_STi.png"),

        "MTI": ("Goulds_3196_PartList_MTi.xlsx", "MTi", "ผังรวม_3196_MTi.png"),

        "LTI": ("Goulds_3196_PartList_LTi.xlsx", "LTi", "ผังรวม_3196_LTi.png"),

        "XLT-I": ("Goulds_3196_PartList_XLT-i.xlsx", "XLT-i", "ผังรวม_3196_XLT-i.png"),

        "XLTI": ("Goulds_3196_PartList_XLT-i.xlsx", "XLT-i", "ผังรวม_3196_XLT-i.png"),

        "I-17": ("Goulds_3196_PartList_i-17.xlsx", "i-17", None),

        "I17": ("Goulds_3196_PartList_i-17.xlsx", "i-17", None),

    }

    

    target_info = None

    for k, info in file_map.items():

        if k == mq_clean or k in mq_clean or mq_clean in k:

            target_info = info

            break

            

    if not target_info:

        target_info = ("Goulds_3196_PartList_STi.xlsx", "STi", "ผังรวม_3196_STi.png")

        

    target_filename, group_name, diag_fname = target_info

    xlsx_path = None

    for gd in goulds_dirs:

        cand = os.path.join(gd, target_filename)

        if os.path.exists(cand):

            xlsx_path = cand

            break

            

    if not xlsx_path:

        return None, f"ไม่พบไฟล์ BOM ({target_filename})"

        

    try:

        wb = openpyxl.load_workbook(xlsx_path, read_only=True)

        ws = wb["Part List"] if "Part List" in wb.sheetnames else wb.active

        rows = list(ws.iter_rows(values_only=True))

    except Exception as e:

        return None, f"Error reading Excel: {e}"

        

    if not rows:

        return None, "BOM Sheet ว่างเปล่า"

        

    parts = []

    for r in rows[1:]:

        if not r or not any(r):

            continue

        item = str(r[0] or "").strip()

        pno = str(r[1] or "").strip()

        desc = str(r[2] or "").strip()

        mat = str(r[3] or "").strip()

        qty = r[4] if len(r) > 4 and r[4] is not None else 1

        price = r[5] if len(r) > 5 and r[5] is not None else None

        parts.append({

            "item": item,

            "part_no": pno if pno != "None" else "",

            "desc": desc,

            "material": mat,

            "qty": qty,

            "price": price,

            "stock_qty": None,

            "stock_status": None

        })

        

    return {

        "ok": True,

        "brand": "GOULDS",

        "model": f"3196 {group_name}",

        "title": f"Goulds 3196 i-FRAME ({group_name})",

        "diagram": f"/api/bom/diagram/{diag_fname}" if diag_fname else None,

        "count": len(parts),

        "parts": parts

    }, None



@app.get("/api/bom/parts")

def get_bom_parts(brand: str = Query("VIKING"), model: str = Query(...)):

    brand_upper = brand.strip().upper()

    if brand_upper == "VIKING":

        res, err = _load_viking_bom_parts(model)

        if not res:

            return {"ok": False, "detail": err or "Not found", "parts": []}

        return res

    elif brand_upper == "GOULDS":

        res, err = _load_goulds_bom_parts(model)

        if not res:

            return {"ok": False, "detail": err or "Not found", "parts": []}

        return res

    else:

        return {"ok": False, "detail": f"ไม่มีฐานข้อมูล BOM สำหรับยี่ห้อ '{brand}' (กรอกเอง)", "parts": []}



@app.get("/api/bom/diagram/{filename}")

def get_bom_diagram(filename: str):

    diag_dirs = [

        os.path.join(DATA_DIR, "bom", "viking", "diagrams"),

        os.path.join(SHARED_VIKING_APP, "diagrams"),

        os.path.join(DATA_DIR, "bom", "goulds", "diagrams"),

        os.path.join(SHARED_GOULDS_DIR, "diagrams")

    ]

    for d in diag_dirs:

        p = os.path.join(d, filename)

        if os.path.exists(p):

            return FileResponse(p, media_type="image/png")

    raise HTTPException(status_code=404, detail="Diagram image not found")



# =============================================================================

# API ROUTES: REPORTS & DRAFTS

# =============================================================================

class SaveReportPayload(BaseModel):

    ticket_no: str

    form_type: Optional[str] = "ALL"

    data: Dict[str, Any]



@app.post("/api/reports/save-draft")

def save_draft(payload: SaveReportPayload):

    ticket = payload.ticket_no.strip()

    if not ticket:

        raise HTTPException(status_code=400, detail="Ticket number required")

    file_path = os.path.join(REPORTS_DIR, f"{ticket}_draft.json")

    payload_data = payload.data

    payload_data["_updated_at"] = datetime.now().isoformat()

    payload_data["_status"] = "draft"

    payload_data["ticket_no"] = ticket

    with open(file_path, "w", encoding="utf-8") as f:

        json.dump(payload_data, f, ensure_ascii=False, indent=2)



    # Pre-compile package so full package PDF is immediately ready for viewing

    full_url = None

    try:

        if payload_data.get("form_instances"):

            sub_res = submit_and_generate(SubmitPayload(ticket_no=ticket, data=payload_data, form_instances=payload_data.get("form_instances")))

            full_url = sub_res.get("generated_files", {}).get("ALL_MERGED")

    except Exception as e:

        print(f"[SaveDraft] Auto compile notice: {e}")



    return {

        "ok": True,

        "ticket_no": ticket,

        "status": "draft",

        "saved_at": payload_data["_updated_at"],

        "full_package_url": full_url or f"/api/download/{ticket}/FULL_PACKAGE_{ticket}.pdf"

    }



@app.get("/api/reports/load-draft/{ticket_no}")

def load_draft(ticket_no: str):

    ticket = ticket_no.strip()

    file_path = os.path.join(REPORTS_DIR, f"{ticket}_draft.json")

    if not os.path.exists(file_path):

        # Also check submitted

        sub_path = os.path.join(REPORTS_DIR, f"{ticket}_submitted.json")

        if os.path.exists(sub_path):

            file_path = sub_path

        else:

            raise HTTPException(status_code=404, detail="No draft found for this ticket")

    try:

        with open(file_path, "r", encoding="utf-8") as f:

            return json.load(f)

    except Exception as e:

        raise HTTPException(status_code=500, detail=str(e))



@app.get("/api/reports")

def list_reports():

    reports = []

    if os.path.exists(REPORTS_DIR):

        for fname in os.listdir(REPORTS_DIR):

            if fname.endswith(".json"):

                fpath = os.path.join(REPORTS_DIR, fname)

                try:

                    with open(fpath, "r", encoding="utf-8") as f:

                        data = json.load(f)

                        reports.append({

                            "ticket_no": data.get("ticket_no", fname.replace(".json", "")),

                            "customer": data.get("customer", "-"),

                            "status": data.get("_status", "saved"),

                            "updated_at": data.get("_updated_at", ""),

                            "serviced_by": data.get("serviced_by", "-"),

                            "generated_files": data.get("_generated_files", [])

                        })

                except:

                    pass

    # Sort by updated_at desc

    reports.sort(key=lambda x: x.get("updated_at", ""), reverse=True)

    return reports



# =============================================================================

# API ROUTES: SUBMIT & ISO PDF GENERATION

# =============================================================================

class SubmitPayload(BaseModel):

    ticket_no: str

    data: Dict[str, Any]

    generate_forms: Optional[List[str]] = None

    form_instances: Optional[List[Dict[str, Any]]] = None



@app.post("/api/reports/submit")

def submit_and_generate(payload: SubmitPayload):

    ticket = payload.ticket_no.strip()

    if not ticket:

        raise HTTPException(status_code=400, detail="Ticket number required")



    data = payload.data

    data["ticket_no"] = ticket

    data["_status"] = "submitted"

    data["_submitted_at"] = datetime.now().isoformat()



    # Create ticket output directory

    ticket_out_dir = os.path.join(OUTPUT_PDF_DIR, ticket)

    os.makedirs(ticket_out_dir, exist_ok=True)



    generated_files = {}

    pdf_paths_for_merge = []



    # Dynamic multi-form package handling

    if payload.form_instances and len(payload.form_instances) > 0:

        # Count frequency of each form code to decide if numbering suffix is needed

        code_counts = {}

        for inst in payload.form_instances:

            c = inst.get("code")

            code_counts[c] = code_counts.get(c, 0) + 1



        code_counters = {}

        for idx, inst in enumerate(payload.form_instances, 1):

            code = inst.get("code")

            code_counters[code] = code_counters.get(code, 0) + 1

            inst_num = inst.get("num") or code_counters[code]

            custom_label = (inst.get("customLabel") or "").strip()



            # Merge base data with instance-specific data

            inst_data = dict(data)

            if inst.get("data") and isinstance(inst.get("data"), dict):

                inst_data.update(inst["data"])

                # Instance / sheet-specific pump identity override

                s_model = inst["data"].get("pump_model") or inst["data"].get("sheet_model")

                if s_model:

                    inst_data["pump_model"] = s_model



                s_serial = inst["data"].get("serial_no") or inst["data"].get("sheet_serial")

                if s_serial:

                    inst_data["serial_no"] = s_serial



                s_code = inst["data"].get("pump_code") or inst["data"].get("sheet_code")

                if s_code:

                    inst_data["pump_code"] = s_code



                s_brand = inst["data"].get("pump_brand") or inst["data"].get("brand")

                if s_brand:

                    inst_data["pump_brand"] = s_brand



            # Ensure Customer, Plant, Contact Person, Tel default to "-" if missing or empty

            for req_field in ["customer", "plant", "contact_person", "contact_no"]:

                f_val = inst_data.get(req_field) or inst_data.get(f"sat_{req_field}") or inst_data.get(f"parts_{req_field}") or data.get(req_field) or data.get(f"sat_{req_field}") or data.get(f"parts_{req_field}")

                inst_data[req_field] = str(f_val).strip() if (f_val and str(f_val).strip() not in ["", "nan", "none"]) else "-"



            if custom_label:

                inst_data["form_instance_label"] = custom_label



            is_multi = code_counts.get(code, 0) > 1

            num_suffix = f"_{inst_num}" if is_multi else ""

            label_text = f" (#{inst_num})" if is_multi else ""

            if custom_label:

                label_text += f" - {custom_label}"



            try:

                if code == "FM-EPS-21":
                    for k in ["customer", "plant", "contact_person", "contact_no", "io_no", "sales_orders", "so"]:
                        v = inst_data.get(k) or inst_data.get(f"exp_{k}")
                        if v:
                            inst_data[k] = v
                    out_file = f"EXP_{ticket}{num_suffix}.pdf"
                    out_path = os.path.join(ticket_out_dir, out_file)
                    pdf_engine.fill_fm_eps_21(inst_data, out_path)
                    generated_files[f"FM-EPS-21{label_text}"] = f"/api/download/{ticket}/{out_file}"
                    pdf_paths_for_merge.append(out_path)

                elif code == "FM-EPS-22":

                    # Custom dates & times for Check Report

                    if inst_data.get("cr_date"):

                        inst_data["date"] = inst_data["cr_date"]

                    if inst_data.get("cr_serviced_date"):

                        inst_data["serviced_date"] = inst_data["cr_serviced_date"]

                    elif inst_data.get("cr_date"):

                        inst_data["serviced_date"] = inst_data["cr_date"]

                    if inst_data.get("cr_time_start"):

                        inst_data["time_start"] = inst_data["cr_time_start"]

                    if inst_data.get("cr_time_finished"):

                        inst_data["time_finished"] = inst_data["cr_time_finished"]

                    if inst_data.get("cr_customer_sign_datetime"):

                        inst_data["customer_sign_datetime"] = inst_data["cr_customer_sign_datetime"]

                    elif inst_data.get("cr_serviced_date") or inst_data.get("cr_date") or inst_data.get("cr_time_finished"):

                        inst_data["customer_sign_datetime"] = None



                    out_file = f"CR_{ticket}{num_suffix}.pdf"

                    out_path = os.path.join(ticket_out_dir, out_file)

                    pdf_engine.fill_fm_eps_22(inst_data, out_path)

                    generated_files[f"FM-EPS-22{label_text}"] = f"/api/download/{ticket}/{out_file}"

                    pdf_paths_for_merge.append(out_path)



                elif code == "FM-EPS-23":

                    out_file = f"SR_{ticket}{num_suffix}.pdf"

                    out_path = os.path.join(ticket_out_dir, out_file)

                    pdf_engine.fill_fm_eps_23(inst_data, out_path)

                    generated_files[f"FM-EPS-23{label_text}"] = f"/api/download/{ticket}/{out_file}"

                    pdf_paths_for_merge.append(out_path)



                elif code == "FM-EPS-25":

                    for k in ["job_no", "io_no", "customer", "plant", "contact_person", "contact_no", "sheet_date", "date"]:

                        v = inst_data.get(k) or inst_data.get(f"parts_{k}")

                        if v:

                            inst_data[k] = v

                    out_file = f"PARTS_{ticket}{num_suffix}.pdf"

                    out_path = os.path.join(ticket_out_dir, out_file)

                    pdf_engine.fill_fm_eps_25(inst_data, out_path)

                    generated_files[f"FM-EPS-25{label_text}"] = f"/api/download/{ticket}/{out_file}"

                    pdf_paths_for_merge.append(out_path)



                elif code == "FM-EPS-31":

                    out_file = f"INSP_{ticket}{num_suffix}.pdf"

                    out_path = os.path.join(ticket_out_dir, out_file)

                    pdf_engine.fill_fm_eps_31(inst_data, out_path)

                    generated_files[f"FM-EPS-31{label_text}"] = f"/api/download/{ticket}/{out_file}"

                    pdf_paths_for_merge.append(out_path)



                elif code == "FM-EPS-45":

                    for k in ["customer", "contact_person", "contact_no", "eval_date"]:

                        v = inst_data.get(f"sat_{k}") or inst_data.get(k)

                        if v:

                            inst_data[k] = v

                    out_file = f"SAT_{ticket}{num_suffix}.pdf"

                    out_path = os.path.join(ticket_out_dir, out_file)

                    pdf_engine.fill_fm_eps_45(inst_data, out_path)

                    generated_files[f"FM-EPS-45{label_text}"] = f"/api/download/{ticket}/{out_file}"

                    pdf_paths_for_merge.append(out_path)



                elif code == "PHOTO_REPORT":

                    out_file = f"PHOTO_{ticket}{num_suffix}.pdf"

                    excel_out_file = f"PHOTO_{ticket}{num_suffix}.xlsx"

                    out_path = os.path.join(ticket_out_dir, out_file)

                    excel_out_path = os.path.join(ticket_out_dir, excel_out_file)



                    # Get photos belonging to this specific instance

                    inst_photos = inst_data.get("photos", [])

                    if not inst_photos and not is_multi:

                        meta_file = os.path.join(PHOTOS_DIR, ticket, "photos.json")

                        if os.path.exists(meta_file):

                            try:

                                with open(meta_file, "r", encoding="utf-8") as pf:

                                    inst_photos = json.load(pf)

                            except:

                                pass



                    pdf_engine.generate_photo_report(inst_data, inst_photos, out_path)

                    try:

                        pdf_engine.generate_photo_report_excel(inst_data, inst_photos, excel_out_path)

                        generated_files[f"PHOTO_REPORT_EXCEL{label_text}"] = f"/api/download/{ticket}/{excel_out_file}"

                    except Exception as ex_err:

                        print(f"Error generating photo excel: {ex_err}")



                    generated_files[f"PHOTO_REPORT{label_text}"] = f"/api/download/{ticket}/{out_file}"

                    pdf_paths_for_merge.append(out_path)



            except Exception as e:

                print(f"Error generating {code} (#{inst_num}): {e}")



    else:

        # Fallback to generate_forms legacy list

        forms_to_gen = payload.generate_forms or ["FM-EPS-22", "FM-EPS-23", "FM-EPS-45"]



        # Ensure Customer, Plant, Contact Person, Tel default to "-" if missing or empty

        for req_field in ["customer", "plant", "contact_person", "contact_no"]:

            f_val = data.get(req_field) or data.get(f"parts_{req_field}")

            data[req_field] = str(f_val).strip() if (f_val and str(f_val).strip() not in ["", "nan", "none"]) else "-"



        # 0. FM-EPS-21 Service Expense
        if "FM-EPS-21" in forms_to_gen:
            exp_path = os.path.join(ticket_out_dir, f"EXP_{ticket}.pdf")
            try:
                for k in ["customer", "plant", "contact_person", "contact_no", "io_no", "sales_orders", "so"]:
                    v = data.get(k) or data.get(f"exp_{k}")
                    if v:
                        data[k] = v
                pdf_engine.fill_fm_eps_21(data, exp_path)
                generated_files["FM-EPS-21"] = f"/api/download/{ticket}/EXP_{ticket}.pdf"
                pdf_paths_for_merge.append(exp_path)
            except Exception as e:
                print(f"Error generating FM-EPS-21: {e}")

        # 1. FM-EPS-22 Check Report

        if "FM-EPS-22" in forms_to_gen:

            cr_path = os.path.join(ticket_out_dir, f"CR_{ticket}.pdf")

            try:

                cr_data = dict(data)

                if cr_data.get("cr_date"):

                    cr_data["date"] = cr_data["cr_date"]

                if cr_data.get("cr_serviced_date"):

                    cr_data["serviced_date"] = cr_data["cr_serviced_date"]

                elif cr_data.get("cr_date"):

                    cr_data["serviced_date"] = cr_data["cr_date"]

                if cr_data.get("cr_time_start"):

                    cr_data["time_start"] = cr_data["cr_time_start"]

                if cr_data.get("cr_time_finished"):

                    cr_data["time_finished"] = cr_data["cr_time_finished"]

                if cr_data.get("cr_customer_sign_datetime"):

                    cr_data["customer_sign_datetime"] = cr_data["cr_customer_sign_datetime"]

                elif cr_data.get("cr_serviced_date") or cr_data.get("cr_date") or cr_data.get("cr_time_finished"):

                    cr_data["customer_sign_datetime"] = None

                pdf_engine.fill_fm_eps_22(cr_data, cr_path)

                generated_files["FM-EPS-22"] = f"/api/download/{ticket}/CR_{ticket}.pdf"

                pdf_paths_for_merge.append(cr_path)

            except Exception as e:

                print(f"Error generating FM-EPS-22: {e}")



        # 2. FM-EPS-23 Service Report

        if "FM-EPS-23" in forms_to_gen:

            sr_path = os.path.join(ticket_out_dir, f"SR_{ticket}.pdf")

            try:

                pdf_engine.fill_fm_eps_23(data, sr_path)

                generated_files["FM-EPS-23"] = f"/api/download/{ticket}/SR_{ticket}.pdf"

                pdf_paths_for_merge.append(sr_path)

            except Exception as e:

                print(f"Error generating FM-EPS-23: {e}")



        # 3. FM-EPS-25 Spare Parts

        if "FM-EPS-25" in forms_to_gen:

            parts_path = os.path.join(ticket_out_dir, f"PARTS_{ticket}.pdf")

            try:

                for k in ["job_no", "io_no", "customer", "plant", "contact_person", "contact_no", "sheet_date", "date"]:

                    v = data.get(f"parts_{k}") or data.get(k)

                    if v:

                        data[k] = v

                pdf_engine.fill_fm_eps_25(data, parts_path)

                generated_files["FM-EPS-25"] = f"/api/download/{ticket}/PARTS_{ticket}.pdf"

                pdf_paths_for_merge.append(parts_path)

            except Exception as e:

                print(f"Error generating FM-EPS-25: {e}")



        # 4. FM-EPS-31 Inspection Report

        if "FM-EPS-31" in forms_to_gen:

            insp_path = os.path.join(ticket_out_dir, f"INSP_{ticket}.pdf")

            try:

                pdf_engine.fill_fm_eps_31(data, insp_path)

                generated_files["FM-EPS-31"] = f"/api/download/{ticket}/INSP_{ticket}.pdf"

                pdf_paths_for_merge.append(insp_path)

            except Exception as e:

                print(f"Error generating FM-EPS-31: {e}")



        # 5. FM-EPS-45 Customer Satisfaction

        if "FM-EPS-45" in forms_to_gen:

            for k in ["customer", "contact_person", "contact_no", "eval_date"]:

                v = data.get(f"sat_{k}") or data.get(k)

                if v:

                    data[k] = v

            sat_path = os.path.join(ticket_out_dir, f"SAT_{ticket}.pdf")

            try:

                pdf_engine.fill_fm_eps_45(data, sat_path)

                generated_files["FM-EPS-45"] = f"/api/download/{ticket}/SAT_{ticket}.pdf"

                pdf_paths_for_merge.append(sat_path)

            except Exception as e:

                print(f"Error generating FM-EPS-45: {e}")



        # 6. Service Photo Report (if photos exist)

        photos = data.get("photos", [])

        if not photos:

            meta_file = os.path.join(PHOTOS_DIR, ticket, "photos.json")

            if os.path.exists(meta_file):

                try:

                    with open(meta_file, "r", encoding="utf-8") as pf:

                        photos = json.load(pf)

                except:

                    pass

        if "PHOTO_REPORT" in forms_to_gen or len(photos) > 0:

            photo_report_path = os.path.join(ticket_out_dir, f"PHOTO_{ticket}.pdf")

            try:

                pdf_engine.generate_photo_report(data, photos, photo_report_path)

                generated_files["PHOTO_REPORT"] = f"/api/download/{ticket}/PHOTO_{ticket}.pdf"

                pdf_paths_for_merge.append(photo_report_path)

            except Exception as e:

                print(f"Error generating PHOTO_REPORT: {e}")



            # Also generate authentic Excel report (.xlsx)

            photo_excel_path = os.path.join(ticket_out_dir, f"PHOTO_{ticket}.xlsx")

            try:

                pdf_engine.generate_photo_report_excel(data, photos, photo_excel_path)

                generated_files["PHOTO_REPORT_EXCEL"] = f"/api/download/{ticket}/PHOTO_{ticket}.xlsx"

            except Exception as e:

                print(f"Error generating PHOTO_REPORT_EXCEL: {e}")



    # 7. Merge All into one complete delivery document (STRICT ISO PACKAGE ORDER)
    # Page 1: CRM Service Request (FM-EPS-28)
    # Page 2: Service Expense (FM-EPS-21, EXP_*) if present
    # Middle Pages: Technical forms (SR, CR, INSP, PARTS, PHOTO)
    # Last Page: Customer Satisfaction Evaluation (FM-EPS-45, SAT_*)
    crm_request_path = pdf_engine.get_or_create_service_request_pdf(ticket, data, ticket_out_dir)
    if crm_request_path and os.path.exists(crm_request_path):
        generated_files["SERVICE_REQUEST"] = f"/api/download/{ticket}/CRM_REQUEST_{ticket}.pdf"

    exp_paths = [p for p in pdf_paths_for_merge if os.path.basename(p).startswith("EXP_")]
    sat_paths = [p for p in pdf_paths_for_merge if os.path.basename(p).startswith("SAT_")]
    other_paths = [p for p in pdf_paths_for_merge if p not in exp_paths and p not in sat_paths]

    ordered_merge_paths = []
    if crm_request_path and os.path.exists(crm_request_path):
        ordered_merge_paths.append(crm_request_path)
    ordered_merge_paths.extend(exp_paths)
    ordered_merge_paths.extend(other_paths)
    ordered_merge_paths.extend(sat_paths)

    if len(ordered_merge_paths) > 1:
        merged_path = os.path.join(ticket_out_dir, f"ALL_{ticket}.pdf")
        try:
            pdf_engine.merge_pdf_package(ordered_merge_paths, merged_path)
            try:
                full_merged_path = os.path.join(ticket_out_dir, f"FULL_PACKAGE_{ticket}.pdf")
                import shutil
                shutil.copyfile(merged_path, full_merged_path)
            except Exception:
                pass
            generated_files["ALL_MERGED"] = f"/api/download/{ticket}/ALL_{ticket}.pdf"
        except Exception as e:
            print(f"Error generating ALL_MERGED: {e}")
    elif len(ordered_merge_paths) == 1:
        merged_path = os.path.join(ticket_out_dir, f"ALL_{ticket}.pdf")
        try:
            import shutil
            shutil.copyfile(ordered_merge_paths[0], merged_path)
            full_merged_path = os.path.join(ticket_out_dir, f"FULL_PACKAGE_{ticket}.pdf")
            shutil.copyfile(merged_path, full_merged_path)
            generated_files["ALL_MERGED"] = f"/api/download/{ticket}/ALL_{ticket}.pdf"
        except Exception:
            pass



    # Save final report state

    data["_generated_files"] = generated_files

    data["_updated_at"] = data["_submitted_at"]

    sub_path = os.path.join(REPORTS_DIR, f"{ticket}_submitted.json")

    with open(sub_path, "w", encoding="utf-8") as f:

        json.dump(data, f, ensure_ascii=False, indent=2)



    # Sync to Google Sheets & Google Drive in background (Hybrid Database)
    try:
        final_pdf_path = None
        if "ALL_MERGED" in generated_files:
            merged_filename = os.path.basename(generated_files["ALL_MERGED"])
            possible_path = os.path.join(ticket_out_dir, merged_filename)
            if os.path.exists(possible_path):
                final_pdf_path = possible_path
        gas_sync.sync_to_google_sheet(ticket, data, pdf_path=final_pdf_path, async_mode=True)
    except Exception as ge:
        print(f"[GAS] Auto-sync notification: {ge}")

    return {

        "ok": True,

        "ticket_no": ticket,

        "status": "ready",

        "generated_files": generated_files

    }



@app.get("/api/download/{ticket_no}/{filename}")

def download_pdf(ticket_no: str, filename: str):

    ticket = ticket_no.strip()

    ticket_out_dir = os.path.join(OUTPUT_PDF_DIR, ticket)

    os.makedirs(ticket_out_dir, exist_ok=True)

    file_path = os.path.join(ticket_out_dir, filename)



    if not os.path.exists(file_path):

        # Auto-generate Photo Report on the fly if requested before package submission

        if filename.startswith("PHOTO_"):

            try:

                t_data = {"ticket_no": ticket}

                # Lookup ticket in CRM cache

                crm_list = load_crm_cache()

                for c_item in crm_list:

                    if str(c_item.get("TicketNo", "")).strip() == ticket:

                        t_data.update({

                            "customer": c_item.get("Customer") or "-",

                            "plant": c_item.get("Site") or "-",

                            "pump_brand": c_item.get("Product") or "-",

                            "model": c_item.get("Model") or "-",

                            "serial_no": c_item.get("SerialNo") or "-",

                            "job_no": c_item.get("SalesOrder") or ticket,

                            "date": c_item.get("AppointmentDate") or c_item.get("CreatedDate") or "-",

                        })

                        break



                # Check saved draft or submitted

                for f_prefix in [f"{ticket}_draft.json", f"{ticket}_submitted.json"]:

                    df_path = os.path.join(REPORTS_DIR, f_prefix)

                    if os.path.exists(df_path):

                        try:

                            with open(df_path, "r", encoding="utf-8") as df:

                                t_data.update(json.load(df))

                        except:

                            pass



                # Load covers

                covers_file = os.path.join(PHOTOS_DIR, ticket, "covers.json")

                if os.path.exists(covers_file):

                    try:

                        with open(covers_file, "r", encoding="utf-8") as cf:

                            t_data["covers"] = json.load(cf)

                    except:

                        pass



                # Load photos

                photos = []

                meta_file = os.path.join(PHOTOS_DIR, ticket, "photos.json")

                if os.path.exists(meta_file):

                    try:

                        with open(meta_file, "r", encoding="utf-8") as pf:

                            photos = json.load(pf)

                    except:

                        pass



                # Load recommended

                rec_file = os.path.join(PHOTOS_DIR, ticket, "recommended.txt")

                if os.path.exists(rec_file):

                    try:

                        with open(rec_file, "r", encoding="utf-8") as rf:

                            t_data["recommended"] = rf.read().strip()

                    except:

                        pass



                pdf_out = os.path.join(ticket_out_dir, f"PHOTO_{ticket}.pdf")

                excel_out = os.path.join(ticket_out_dir, f"PHOTO_{ticket}.xlsx")



                pdf_engine.generate_photo_report(t_data, photos, pdf_out)

                try:

                    pdf_engine.generate_photo_report_excel(t_data, photos, excel_out)

                except Exception as ex:

                    print(f"[AutoGenerate] Excel export error: {ex}")

            except Exception as e:

                print(f"[AutoGenerate] Error on-demand generating {filename}: {e}")



    if not os.path.exists(file_path):

        raise HTTPException(status_code=404, detail="Requested file not found")

    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if filename.endswith(".xlsx") else "application/pdf"

    return FileResponse(file_path, media_type=media_type, filename=filename)



@app.post("/api/photos/{ticket_no}/generate")

async def generate_single_photo_report(ticket_no: str, payload: Dict[str, Any] = Body(...)):

    ticket = ticket_no.strip()

    ticket_out_dir = os.path.join(OUTPUT_PDF_DIR, ticket)

    os.makedirs(ticket_out_dir, exist_ok=True)



    photos = payload.get("photos", [])

    if not photos:

        meta_file = os.path.join(PHOTOS_DIR, ticket, "photos.json")

        if os.path.exists(meta_file):

            try:

                with open(meta_file, "r", encoding="utf-8") as pf:

                    photos = json.load(pf)

            except:

                pass



    if "covers" not in payload:

        covers_file = os.path.join(PHOTOS_DIR, ticket, "covers.json")

        if os.path.exists(covers_file):

            try:

                with open(covers_file, "r", encoding="utf-8") as cf:

                    payload["covers"] = json.load(cf)

            except:

                pass



    pdf_out = os.path.join(ticket_out_dir, f"PHOTO_{ticket}.pdf")

    excel_out = os.path.join(ticket_out_dir, f"PHOTO_{ticket}.xlsx")



    pdf_engine.generate_photo_report(payload, photos, pdf_out)

    excel_ok = False

    try:

        pdf_engine.generate_photo_report_excel(payload, photos, excel_out)

        excel_ok = True

    except Exception as e:

        print(f"Error generating photo excel: {e}")



    return {

        "ok": True,

        "pdf_url": f"/api/download/{ticket}/PHOTO_{ticket}.pdf",

        "excel_url": f"/api/download/{ticket}/PHOTO_{ticket}.xlsx" if excel_ok else None

    }



@app.post("/api/photos/{ticket_no}/update_meta")

async def update_photo_meta(ticket_no: str, payload: Dict[str, Any] = Body(...)):

    ticket = ticket_no.strip()

    ticket_photos_dir = os.path.join(PHOTOS_DIR, ticket)

    os.makedirs(ticket_photos_dir, exist_ok=True)

    meta_file = os.path.join(ticket_photos_dir, "photos.json")



    updated_photos = payload.get("photos", [])

    with open(meta_file, "w", encoding="utf-8") as f:

        json.dump(updated_photos, f, ensure_ascii=False, indent=2)



    # Save covers if present

    if "covers" in payload:

        covers_file = os.path.join(ticket_photos_dir, "covers.json")

        with open(covers_file, "w", encoding="utf-8") as f:

            json.dump(payload["covers"], f, ensure_ascii=False, indent=2)



    # Save recommended text if present

    if "recommended" in payload:

        rec_file = os.path.join(ticket_photos_dir, "recommended.txt")

        with open(rec_file, "w", encoding="utf-8") as f:

            f.write(str(payload.get("recommended") or ""))



    return {"ok": True}



@app.get("/api/photos/{ticket_no}/covers")

def get_photo_covers(ticket_no: str):

    ticket = ticket_no.strip()

    covers_file = os.path.join(PHOTOS_DIR, ticket, "covers.json")

    if os.path.exists(covers_file):

        try:

            with open(covers_file, "r", encoding="utf-8") as cf:

                return json.load(cf)

        except:

            pass

    return []



@app.get("/api/photos/{ticket_no}/recommended")

def get_photo_recommended(ticket_no: str):

    ticket = ticket_no.strip()

    rec_file = os.path.join(PHOTOS_DIR, ticket, "recommended.txt")

    if os.path.exists(rec_file):

        try:

            with open(rec_file, "r", encoding="utf-8") as rf:

                return {"recommended": rf.read().strip()}

        except:

            pass

    return {"recommended": ""}



# =============================================================================

# API ROUTES: SERVICE PHOTO MANAGEMENT

# =============================================================================

@app.post("/api/photos/upload")

async def upload_photo(

    ticket_no: str = Form(...),

    category: str = Form("General"),

    caption: str = Form(""),

    cover_id: Optional[str] = Form(None),

    file: UploadFile = File(...)

):

    ticket = ticket_no.strip()

    ticket_photos_dir = os.path.join(PHOTOS_DIR, ticket)

    os.makedirs(ticket_photos_dir, exist_ok=True)



    file_ext = os.path.splitext(file.filename)[1] or ".jpg"

    photo_id = f"photo_{int(datetime.now().timestamp())}_{uuid.uuid4().hex[:6]}"

    saved_filename = f"{photo_id}{file_ext}"

    saved_path = os.path.join(ticket_photos_dir, saved_filename)



    with open(saved_path, "wb") as buffer:

        shutil.copyfileobj(file.file, buffer)



    # Read/Update photos metadata

    meta_file = os.path.join(ticket_photos_dir, "photos.json")

    photos_meta = []

    if os.path.exists(meta_file):

        try:

            with open(meta_file, "r", encoding="utf-8") as f:

                photos_meta = json.load(f)

        except:

            pass



    photo_entry = {

        "id": photo_id,

        "filename": saved_filename,

        "category": category,

        "caption": caption,

        "cover_id": cover_id or "cover_1",

        "uploaded_at": datetime.now().isoformat(),

        "url": f"/api/photos/view/{ticket}/{saved_filename}",

        "file_path": saved_path

    }

    photos_meta.append(photo_entry)



    with open(meta_file, "w", encoding="utf-8") as f:

        json.dump(photos_meta, f, ensure_ascii=False, indent=2)



    return {"ok": True, "photo": photo_entry}



@app.get("/api/photos/{ticket_no}")

def get_ticket_photos(ticket_no: str):

    ticket = ticket_no.strip()

    meta_file = os.path.join(PHOTOS_DIR, ticket, "photos.json")

    if os.path.exists(meta_file):

        try:

            with open(meta_file, "r", encoding="utf-8") as f:

                return json.load(f)

        except:

            pass

    return []



@app.get("/api/photos/view/{ticket_no}/{filename}")

def view_photo(ticket_no: str, filename: str):

    file_path = os.path.join(PHOTOS_DIR, ticket_no, filename)

    if not os.path.exists(file_path):

        raise HTTPException(status_code=404, detail="Photo not found")

    return FileResponse(file_path)



@app.delete("/api/photos/{ticket_no}/{photo_id}")

def delete_photo(ticket_no: str, photo_id: str):

    ticket = ticket_no.strip()

    ticket_photos_dir = os.path.join(PHOTOS_DIR, ticket)

    meta_file = os.path.join(ticket_photos_dir, "photos.json")

    if os.path.exists(meta_file):

        try:

            with open(meta_file, "r", encoding="utf-8") as f:

                photos_meta = json.load(f)

            remaining = []

            deleted = False

            for p in photos_meta:

                if p.get("id") == photo_id:

                    # Remove file from disk

                    fpath = os.path.join(ticket_photos_dir, p.get("filename", ""))

                    if os.path.exists(fpath):

                        os.remove(fpath)

                    deleted = True

                else:

                    remaining.append(p)

            with open(meta_file, "w", encoding="utf-8") as f:

                json.dump(remaining, f, ensure_ascii=False, indent=2)

            return {"ok": True, "deleted": deleted}

        except Exception as e:

            raise HTTPException(status_code=500, detail=str(e))

    return {"ok": False, "message": "No photos found"}



# =============================================================================
# API ROUTES: GOOGLE APPS SCRIPT & GOOGLE SHEETS SYNC
# =============================================================================

class GasConfigPayload(BaseModel):
    webapp_url: str
    enabled: Optional[bool] = True

@app.get("/api/gas/config")
def get_gas_status():
    return gas_sync.get_gas_config()

@app.post("/api/gas/config")
def update_gas_status(payload: GasConfigPayload):
    saved = gas_sync.save_gas_config(payload.webapp_url, payload.enabled)
    return {"ok": True, "config": saved}

@app.post("/api/gas/test")
def test_gas_endpoint(payload: GasConfigPayload):
    res = gas_sync.test_gas_connection(payload.webapp_url)
    return res

@app.post("/api/gas/sync/{ticket_no}")
def manual_gas_sync(ticket_no: str):
    ticket = ticket_no.strip()
    file_path = os.path.join(REPORTS_DIR, f"{ticket}_submitted.json")
    if not os.path.exists(file_path):
        file_path = os.path.join(REPORTS_DIR, f"{ticket}_draft.json")
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Ticket data not found")
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    pdf_path = os.path.join(OUTPUT_PDF_DIR, ticket, f"ALL_{ticket}.pdf")
    if not os.path.exists(pdf_path):
        pdf_path = os.path.join(OUTPUT_PDF_DIR, ticket, f"FULL_PACKAGE_{ticket}.pdf")
    res = gas_sync.sync_to_google_sheet(ticket, data, pdf_path=pdf_path if os.path.exists(pdf_path) else None, async_mode=False)
    return res




# =============================================================================

# MOUNT STATIC FRONTEND FILES

# =============================================================================

if os.path.exists(FRONTEND_DIR):

    app.mount("/", StaticFiles(directory=FRONTEND_DIR, html=True), name="static")



if __name__ == "__main__":

    import uvicorn

    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=True)



